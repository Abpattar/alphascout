"""Professional formatting pass on Alpha_Scout_Final_Report.docx"""
import copy
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement as _OE

PATH = '/home/neo/Codes/alphascout/Alpha_Scout_Final_Report.docx'
doc = docx.Document(PATH)
paras = doc.paragraphs

def set_page_break_before(par):
    par.paragraph_format.page_break_before = True

def delete_paragraph(par):
    el = par._p
    el.getparent().remove(el)

# ---------------------------------------------------------------
# 1. Front-matter headings: unify formatting (bold 16pt, centered, 12/12)
# ---------------------------------------------------------------
for idx in [51, 63, 96, 110, 127]:  # DECLARATION, CERT, ACKNOWLEDGEMENTS, TOC, ABSTRACT
    p = paras[idx]
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.keep_with_next = True
    for r in p.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.name = 'Times New Roman'

# ---------------------------------------------------------------
# 2. Completion certificate section: replace paras 64-95 with a real certificate
# ---------------------------------------------------------------
cert_heading = paras[63]
empty_cert_paras = paras[64:96]

text = ('This is to certify that Mr. Aditya B Pattar (USN: 20241CIT0115), a student of '
        'B.Tech in CIT-Computer Science & Engineering (Internet of Things), Presidency School '
        'of Artificial Intelligence & Advanced Computing, Presidency University, Bengaluru, has '
        'successfully completed the internship work titled \u201cAlphaScout v1.0: Multi-Sector '
        'Small-Cap News to Trade Signal Bot for Indian Markets\u201d during the academic year '
        '2026-2027, in partial fulfilment of the requirements for the CSS7000 Internship '
        'course. The work was carried out under the guidance of Mr. Francis Annareddy, '
        'Assistant Professor, Program Internship Coordinator, Presidency School of Artificial '
        'Intelligence & Advanced Computing, Presidency University.')

tbl = doc.add_table(rows=2, cols=2)
tbl.style = doc.styles['Table Grid']
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = True

# Row 0: merged certificate text
r0 = tbl.rows[0]
c00 = r0.cells[0]
c01 = r0.cells[1]
c00.merge(c01)
p = c00.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(6)
run = p.add_run('This is to certify that ')
run.font.name = 'Times New Roman'
run.font.size = Pt(12)
run.font.bold = True
run2 = p.add_run(text[len('This is to certify that '):])
run2.font.name = 'Times New Roman'
run2.font.size = Pt(12)

# Row 1: guide / HOD signature blocks
r1 = tbl.rows[1]
guide = r1.cells[0]
hod = r1.cells[1]
for cell, lines in ((guide, ['Mr. Francis Annareddy', 'Assistant Professor',
                             'Internship Guide', 'Presidency University']),
                    (hod, ['Dr. Anandaraj S.P', 'Professor & Head of the Department',
                           'PSAIAC', 'Presidency University'])):
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    for j, ln in enumerate(lines):
        r = cp.add_run(ln)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = (j == 0)
        if j < len(lines) - 1:
            r.add_break()

# Move the table right after the certificate heading
cert_heading._p.addnext(tbl._tbl)

# Delete the 32 empty paragraphs
for p in empty_cert_paras:
    delete_paragraph(p)

# ---------------------------------------------------------------
# 3. Page breaks before ACKNOWLEDGEMENTS / TABLE OF CONTENTS / ABSTRACT / CHAPTER 1
# ---------------------------------------------------------------
set_page_break_before(paras[96])   # ACKNOWLEDGEMENTS
set_page_break_before(paras[110])  # TABLE OF CONTENTS
set_page_break_before(paras[127])  # ABSTRACT

# Trim spacers after acknowledgement signature (105-109 -> 2 kept) and after TOC table (111-126 -> 0)
keep_after_ack = 2
keep_after_toc = 0

to_delete_ack = paras[105:109]
for p in to_delete_ack[:5 - keep_after_ack]:
    delete_paragraph(p)

ch1 = None
to_delete_toc = []
for p in paras[111:127]:
    if p.text.strip():
        continue
    to_delete_toc.append(p)
for p in to_delete_toc[keep_after_toc:]:
    delete_paragraph(p)

# Abstract tail spacers (134-156) -> delete all, then page-break before CHAPTER 1
to_delete_abs = [p for p in paras[134:157] if not p.text.strip()]
for p in to_delete_abs:
    delete_paragraph(p)
set_page_break_before(paras[157])  # CHAPTER 1

# ---------------------------------------------------------------
# 4. Heading styles: keep with next + keep lines
# ---------------------------------------------------------------
for sname in ['Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[sname]
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True

# ---------------------------------------------------------------
# 5. Body text: unify spacing (6pt after) for justified 1.5-spaced Normal paragraphs
# ---------------------------------------------------------------
for p in doc.paragraphs:
    if not p.text.strip():
        continue
    pf = p.paragraph_format
    if pf.line_spacing != 1.5:
        continue
    pPr = p._p.pPr
    if pPr is None:
        continue
    jc = pPr.find(qn('w:jc'))
    if jc is None or jc.get(qn('w:val')) != 'both':
        continue
    pf.space_after = Pt(6)

# ---------------------------------------------------------------
# 6. Tables: cantSplit, header repeat, keep together for small tables
# ---------------------------------------------------------------
for ti in range(4, 14):
    tbl = doc.tables[ti]
    nrows = len(tbl.rows)
    for ri, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(_OE('w:cantSplit'))
        if nrows > 8 and ri == 0:
            if trPr.find(qn('w:tblHeader')) is None:
                trPr.append(_OE('w:tblHeader'))
    if nrows <= 8:
        for row in tbl.rows[:-1]:
            last_cell = row.cells[-1]
            last_par = last_cell.paragraphs[-1]
            last_par.paragraph_format.keep_with_next = True

doc.save(PATH)
print('OK - saved')